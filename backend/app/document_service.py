"""Extract readable course text without Canvas requests, storage, or AI calls."""

import hashlib
import io
import json
import re
from html.parser import HTMLParser
from zipfile import BadZipFile, ZipFile

import pypdf

MAX_PDF_PAGES = 1000
MAX_PPTX_SLIDES = 1000
MAX_OFFICE_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_OFFICE_PARTS = 10000
PIPELINE_VERSION = f"course-text-v1:pypdf-{pypdf.__version__}"


class _TextParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self.hidden += 1
        if tag in ("p", "div", "li", "br", "h1", "h2", "h3", "h4", "h5", "h6", "tr"):
            self.parts.append("\n")
        elif tag in ("td", "th"):
            self.parts.append("\t")
        if tag == "img" and not self.hidden:
            alt = dict(attrs).get("alt")
            if alt:
                self.parts.append(f" {alt} ")

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self.hidden = max(0, self.hidden - 1)
        if tag in ("p", "div", "li", "h1", "h2", "h3", "h4", "h5", "h6", "tr"):
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def _clean_text(text):
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.rstrip() for line in text.split("\n")]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _result(content, format_name, extractor_version):
    return {
        "status": "no_text", "format": format_name, "text": "", "pages": [],
        "page_count": None, "character_count": 0,
        "content_hash": hashlib.sha256(content).hexdigest(),
        "extractor_version": extractor_version,
    }


def extract_html(html):
    """Turn API-provided HTML into plain text; never visit its linked pages."""
    result = _result(html.encode("utf-8"), "html", "html-text-v2")
    parser = _TextParser()
    parser.feed(html)
    result["text"] = _clean_text("".join(parser.parts))
    result["character_count"] = len(result["text"])
    result["status"] = "ok" if result["text"] else "no_text"
    return result


def extract_pdf(content):
    """Read a PDF's text layer and retain page numbers and incomplete pages."""
    result = _result(content, "pdf", f"pypdf-{pypdf.__version__}:text-v1")
    if not content.lstrip().startswith(b"%PDF-"):
        return {**result, "status": "error", "error_type": "InvalidPDF"}
    try:
        reader = pypdf.PdfReader(io.BytesIO(content))
        if reader.is_encrypted and not reader.decrypt(""):
            return {**result, "status": "encrypted"}
        result["page_count"] = len(reader.pages)
        if result["page_count"] > MAX_PDF_PAGES:
            return {**result, "status": "too_large"}
        for number, page in enumerate(reader.pages, start=1):
            try:
                text = _clean_text(page.extract_text() or "")
                result["pages"].append({"page_number": number, "text": text, "status": "ok" if text else "no_text"})
            except Exception as error:
                result["pages"].append({"page_number": number, "text": "", "status": "error", "error_type": type(error).__name__})
        result["text"] = "\n\n".join(page["text"] for page in result["pages"] if page["text"])
        result["character_count"] = len(result["text"])
        if result["text"]:
            result["status"] = "ok" if all(page["status"] == "ok" for page in result["pages"]) else "partial"
        elif any(page["status"] == "error" for page in result["pages"]):
            result["status"] = "error"
    except Exception as error:
        result.update(status="error", error_type=type(error).__name__)
    return result


def file_format(metadata):
    """Identify supported formats before spending bandwidth on a download."""
    content_type = (metadata.get("content-type") or metadata.get("content_type") or "").split(";")[0].strip().lower()
    filename = (metadata.get("filename") or metadata.get("display_name") or "").lower()
    if content_type == "application/pdf" or filename.endswith(".pdf"):
        return "pdf"
    if content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document" or filename.endswith(".docx"):
        return "docx"
    if content_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation" or filename.endswith(".pptx"):
        return "pptx"
    if content_type in ("application/msword", "application/vnd.ms-word") or filename.endswith(".doc"):
        return "doc"
    if content_type in ("application/vnd.ms-powerpoint", "application/mspowerpoint", "application/x-mspowerpoint") or filename.endswith(".ppt"):
        return "ppt"
    if content_type == "application/x-ipynb+json" or filename.endswith(".ipynb"):
        return "notebook"
    if content_type == "text/html" or filename.endswith((".html", ".htm")):
        return "html"
    if content_type.startswith("text/") or filename.endswith((".txt", ".md", ".csv")):
        return "text"
    return None


def file_pipeline_version(metadata):
    """Invalidate Office caches when their reader changes, without rereading PDFs."""
    format_name = file_format(metadata)
    if format_name == "docx":
        import docx
        return f"{PIPELINE_VERSION}:docx-{docx.__version__}:text-v1"
    if format_name == "pptx":
        import pptx
        return f"{PIPELINE_VERSION}:pptx-{pptx.__version__}:text-v1"
    if format_name in ("doc", "ppt"):
        from backend.app.legacy_office import converter_fingerprint
        modern = "docx" if format_name == "doc" else "pptx"
        reader_version = file_pipeline_version({"filename": f"input.{modern}"})
        return f"{reader_version}:legacy-v2:{converter_fingerprint(format_name)}"
    return PIPELINE_VERSION


def _office_package_issue(content, main_part):
    """Check Office ZIP size before a parser expands all its contents in memory."""
    try:
        with ZipFile(io.BytesIO(content)) as package:
            parts = package.infolist()
            if len(parts) > MAX_OFFICE_PARTS or sum(part.file_size for part in parts) > MAX_OFFICE_UNCOMPRESSED_BYTES:
                return {"status": "too_large", "error_type": "OfficePackageLimit"}
            if any(part.flag_bits & 1 for part in parts):
                return {"status": "encrypted", "error_type": "EncryptedOfficeArchive"}
            if main_part not in package.namelist():
                return {"status": "error", "error_type": "WrongOfficeFormat"}
    except (BadZipFile, ValueError):
        return {"status": "error", "error_type": "InvalidOfficeArchive"}
    return None


def _word_block_text(block, depth=0):
    """Read paragraphs and nested tables in order, without repeating merged cells."""
    from docx.text.paragraph import Paragraph

    if isinstance(block, Paragraph):
        return _clean_text(block.text)
    if depth >= 20:
        raise ValueError("Nested table limit")
    rows, seen_cells = [], set()
    for row in block.rows:
        values = []
        for cell in row.cells:
            # Merged grid positions share one underlying XML cell.
            if cell._tc in seen_cells:
                values.append("")
                continue
            seen_cells.add(cell._tc)
            values.append("\n".join(_word_block_text(child, depth + 1) for child in cell.iter_inner_content()).strip())
        rows.append("\t".join(values))
    return "\n".join(rows).strip("\n")


def extract_docx(content):
    """Read Word body/header/footer text with block references, not invented pages."""
    import docx
    from docx.text.paragraph import Paragraph

    result = _result(content, "docx", f"python-docx-{docx.__version__}:text-v1")
    result["blocks"] = []
    issue = _office_package_issue(content, "word/document.xml")
    if issue:
        return {**result, **issue}
    try:
        document = docx.Document(io.BytesIO(content))
        containers = [("body", document)]
        seen_parts = set()
        for number, section in enumerate(document.sections, start=1):
            for name in ("header", "first_page_header", "even_page_header", "footer", "first_page_footer", "even_page_footer"):
                container = getattr(section, name)
                if container.is_linked_to_previous or container.part.partname in seen_parts:
                    continue
                seen_parts.add(container.part.partname)
                containers.append((f"section {number} {name}", container))
        warnings = set()
        unread_tags = {
            "drawing": "images_not_extracted", "pict": "images_not_extracted",
            "object": "embedded_content_not_extracted", "oMath": "equations_not_extracted",
            "footnoteReference": "footnotes_not_extracted", "endnoteReference": "footnotes_not_extracted",
            "txbxContent": "text_boxes_not_extracted", "sdt": "content_controls_not_extracted",
            "ins": "tracked_changes_not_extracted", "del": "tracked_changes_not_extracted",
        }
        for location, container in containers:
            for element in container._element.iter():
                tag = element.tag.rsplit("}", 1)[-1] if isinstance(element.tag, str) else ""
                if tag in unread_tags:
                    warnings.add(unread_tags[tag])
            for number, block in enumerate(container.iter_inner_content(), start=1):
                entry = {"location": location, "block_number": number, "kind": "paragraph" if isinstance(block, Paragraph) else "table"}
                try:
                    entry.update(text=_word_block_text(block), status="ok")
                except Exception as error:
                    entry.update(text="", status="error", error_type=type(error).__name__)
                    warnings.add("block_read_error")
                if entry["text"] or entry["status"] == "error":
                    result["blocks"].append(entry)
        result["text"] = "\n\n".join(block["text"] for block in result["blocks"] if block["text"])
        result["block_count"] = len(result["blocks"])
        result["character_count"] = len(result["text"])
        result["warnings"] = sorted(warnings)
        if result["text"]:
            result["status"] = "partial" if warnings else "ok"
        elif "block_read_error" in warnings:
            result["status"] = "error"
    except Exception as error:
        result.update(status="error", error_type=type(error).__name__)
    return result


def _slide_shape_text(shapes, depth=0):
    """Read grouped text and tables; flag visuals instead of pretending to read them."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    texts, warnings = [], set()
    for shape in shapes:
        try:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                if depth >= 20:
                    warnings.add("group_depth_limit")
                    continue
                child_text, child_warnings = _slide_shape_text(shape.shapes, depth + 1)
                texts.extend(child_text)
                warnings.update(child_warnings)
            elif shape.has_text_frame:
                text = "\n".join("  " * paragraph.level + paragraph.text.replace("\v", "\n") for paragraph in shape.text_frame.paragraphs)
                if text.strip():
                    texts.append(_clean_text(text))
            elif shape.has_table:
                rows = ["\t".join("" if cell.is_spanned else _clean_text(cell.text) for cell in row.cells) for row in shape.table.rows]
                text = "\n".join(rows).strip("\n")
                if text.strip():
                    texts.append(text)
            elif shape.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE):
                warnings.add("images_not_extracted")
            elif shape.shape_type != MSO_SHAPE_TYPE.LINE:
                warnings.add("drawing_or_embedded_content_not_extracted")
        except Exception:
            warnings.add("shape_read_error")
    return texts, warnings


def extract_pptx(content):
    """Read slide text/tables and speaker notes while preserving slide numbers."""
    import pptx

    result = _result(content, "pptx", f"python-pptx-{pptx.__version__}:text-v1")
    result["slides"] = []
    issue = _office_package_issue(content, "ppt/presentation.xml")
    if issue:
        return {**result, **issue}
    try:
        presentation = pptx.Presentation(io.BytesIO(content))
        result["slide_count"] = len(presentation.slides)
        if result["slide_count"] > MAX_PPTX_SLIDES:
            return {**result, "status": "too_large"}
        for number, slide in enumerate(presentation.slides, start=1):
            entry = {"slide_number": number, "text": "", "status": "no_text"}
            try:
                texts, warnings = _slide_shape_text(slide.shapes)
                for element in slide._element.iter():
                    tag = element.tag.rsplit("}", 1)[-1] if isinstance(element.tag, str) else ""
                    if tag == "oMath":
                        warnings.add("equations_not_extracted")
                    elif tag == "blip":
                        warnings.add("images_not_extracted")
                title = slide.shapes.title
                entry["title"] = _clean_text(title.text.replace("\v", "\n")) if title is not None else None
                notes = ""
                if slide.has_notes_slide:
                    frame = slide.notes_slide.notes_text_frame
                    notes = _clean_text(frame.text.replace("\v", "\n")) if frame is not None else ""
                text = "\n\n".join(texts)
                if notes:
                    prefix = "\n\nSpeaker notes:\n" if text else "Speaker notes:\n"
                    entry["notes_start_character"] = len(text) + len(prefix)
                    text += prefix + notes
                entry.update(text=text, warnings=sorted(warnings), status=("partial" if warnings else "ok") if text else "no_text")
            except Exception as error:
                entry.update(status="error", error_type=type(error).__name__)
            result["slides"].append(entry)
        result["text"] = "\n\n".join(slide["text"] for slide in result["slides"] if slide["text"])
        result["character_count"] = len(result["text"])
        if result["text"]:
            result["status"] = "ok" if all(slide["status"] == "ok" for slide in result["slides"]) else "partial"
        elif any(slide["status"] == "error" for slide in result["slides"]):
            result["status"] = "error"
    except Exception as error:
        result.update(status="error", error_type=type(error).__name__)
    return result


def extract_notebook(content):
    """Read notebook instructions and code as text, without executing any cells.

    Character offsets retain cell provenance. Outputs and embedded images are
    not extracted; notebooks requiring them are explicitly marked partial.
    """
    result = _result(content, "notebook", "notebook-text-v1")
    result["cells"] = []
    try:
        notebook = json.loads(content.decode("utf-8-sig"))
        if not isinstance(notebook, dict) or notebook.get("nbformat") != 4 or not isinstance(notebook.get("cells"), list):
            raise ValueError("Unsupported notebook structure")
        parts = []
        incomplete = False
        offset = 0
        for number, cell in enumerate(notebook["cells"], start=1):
            if not isinstance(cell, dict) or cell.get("cell_type") not in ("markdown", "code", "raw"):
                incomplete = True
                continue
            source = cell.get("source", "")
            if isinstance(source, list) and all(isinstance(line, str) for line in source):
                source = "".join(source)
            if not isinstance(source, str):
                incomplete = True
                continue
            if cell.get("attachments") or cell.get("outputs"):
                incomplete = True
            text = source.replace("\r\n", "\n").replace("\r", "\n").strip("\n") if cell["cell_type"] == "code" else _clean_text(source)
            if cell["cell_type"] == "markdown" and re.search(r"!\[[^\]]*\]\(", source):
                incomplete = True
            if text:
                text = f"[{cell['cell_type']} cell {number}]\n{text}"
                start = offset + (2 if parts else 0)
                result["cells"].append({"cell_number": number, "cell_type": cell["cell_type"], "start_character": start, "end_character": start + len(text)})
                parts.append(text)
                offset = start + len(text)
        result["text"] = "\n\n".join(parts)
        result["character_count"] = len(result["text"])
        result["status"] = ("partial" if incomplete else "ok") if result["text"] else "no_text"
    except (UnicodeDecodeError, json.JSONDecodeError, ValueError):
        result.update(status="error", error_type="InvalidNotebook")
    return result


def extract_legacy_office(content, format_name):
    """Convert legacy bytes, reuse modern readers, and keep the original file hash."""
    from backend.app.legacy_office import LegacyConversionError, convert_legacy_office, converter_fingerprint

    version = f"legacy-v2:{converter_fingerprint(format_name)}"
    original = _result(content, format_name, version)
    try:
        converted = convert_legacy_office(content, format_name)
    except LegacyConversionError as error:
        return {**original, "status": error.status, "error_type": error.error_type}
    modern = "docx" if format_name == "doc" else "pptx"
    result = extract_docx(converted) if modern == "docx" else extract_pptx(converted)
    if format_name == "doc":
        # Legacy import can change document structure. macOS has a table bridge,
        # but converted Word layout still should not be treated as verified.
        result["warnings"] = sorted(set(result.get("warnings", [])) | {"legacy_doc_structure_unverified"})
        if result.get("text") and result["status"] == "ok":
            result["status"] = "partial"
    return {
        **result, "format": format_name, "converted_format": modern,
        "content_hash": original["content_hash"],
        "extractor_version": f"{version}:{result['extractor_version']}",
    }


def extract_file(metadata, content):
    """Select a supported reader and explicitly flag other file formats."""
    format_name = file_format(metadata)
    if format_name == "pdf":
        return extract_pdf(content)
    if format_name == "docx":
        return extract_docx(content)
    if format_name == "pptx":
        return extract_pptx(content)
    if format_name in ("doc", "ppt"):
        return extract_legacy_office(content, format_name)
    if format_name == "notebook":
        return extract_notebook(content)
    if format_name in ("text", "html"):
        try:
            encoding = "utf-16" if content.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
            text = content.decode(encoding)
        except UnicodeDecodeError:
            return {**_result(content, "text", "text-v1"), "status": "error", "error_type": "UnsupportedTextEncoding"}
        if format_name == "html":
            result = extract_html(text)
            result["content_hash"] = hashlib.sha256(content).hexdigest()
            return result
        result = _result(content, "text", "text-v1")
        result["text"] = _clean_text(text)
        result["character_count"] = len(result["text"])
        result["status"] = "ok" if result["text"] else "no_text"
        return result
    return {**_result(content, "unsupported", "none"), "status": "unsupported"}
