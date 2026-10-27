import copy
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

import requests

from backend.app import ingestion_service
from backend.app.ingestion_store import IngestionStore
from backend.tests.test_document_service import make_pdf


def record(canvas_id, source_type, **metadata):
    return {
        "canvas_id": canvas_id, "course_id": 42, "source_type": source_type,
        "title": metadata.get("title") or metadata.get("name") or "Course material",
        "url": "https://canvas.example.edu/courses/42", "metadata": metadata,
    }


def snapshot():
    result = {
        "course_id": 42, "synced_at": "first",
        "course": record(42, "course", course_code="CSE 1000"),
        "syllabus": record(42, "syllabus", syllabus_body="<p>Late work policy</p>"),
        "source_status": {source: {"status": "ok"} for source in (
            "course", "syllabus", "assignments", "modules", "module_items", "pages", "files", "quizzes", "announcements", "discussions",
        )},
        "file_inventory": [{
            "course_id": 42, "file_id": 9, "title": "Lecture", "metadata": None,
            "references": [{"source_type": "module_item", "canvas_id": 100, "title": "Lecture", "module_id": 5, "module_name": "Week 1"}],
        }],
    }
    result.update({source: [] for source in ("assignments", "modules", "module_items", "pages", "files", "quizzes", "announcements", "discussions")})
    result["assignments"] = [record(7, "assignment", description="<p>Read the lecture</p>")]
    return result


class IngestionPipelineTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.store = IngestionStore("test-student", Path(self.directory.name) / "ingestion.sqlite3")
        self.metadata = {"id": 9, "filename": "lecture.pdf", "display_name": "Lecture.pdf", "updated_at": "2026-10-01T00:00:00Z", "size": 1024, "url": "https://canvas.example.edu/files/9/download"}
        self.content = make_pdf(["Linear algebra lecture"])
        self.snapshot = snapshot()
        self.discover = patch.object(ingestion_service, "sync_course", side_effect=lambda *args, **kwargs: copy.deepcopy(self.snapshot)).start()
        self.info = patch.object(ingestion_service.canvas_service, "get_file_info", side_effect=lambda *args: copy.deepcopy(self.metadata)).start()
        self.download = patch.object(ingestion_service.canvas_service, "download_canvas_file", side_effect=lambda *args, **kwargs: (kwargs["metadata"], self.content)).start()
        self.addCleanup(patch.stopall)
        self.addCleanup(self.directory.cleanup)

    def run_ingestion(self, **kwargs):
        return ingestion_service.ingest_course(42, store=self.store, **kwargs)

    def test_first_run_extracts_native_html_and_pdf_and_keeps_provenance(self):
        report = self.run_ingestion()
        self.assertEqual(report["statistics"]["downloads"], 1)
        self.assertEqual(report["statistics"]["readable"], 3)
        self.assertEqual(report["snapshot"]["coverage"], "ok")
        documents = self.store.load_documents(42, readable_only=True)
        self.assertEqual(len(documents), 3)
        pdf = self.store.load_document(42, "file:9")
        self.assertIn("Linear algebra lecture", pdf["text"])
        self.assertEqual(pdf["references"][0]["module_id"], 5)
        self.assertEqual(pdf["pages"][0]["page_number"], 1)
        self.discover.assert_called_once_with(42, include_page_bodies=True)
        self.info.assert_called_once_with(42, 9)
        self.assertEqual(self.download.call_args.kwargs["metadata"]["id"], 9)

    def test_second_run_uses_fresh_metadata_but_never_redownloads_unchanged_pdf(self):
        self.run_ingestion()
        self.snapshot["synced_at"] = "second"
        report = self.run_ingestion()
        self.assertEqual(report["statistics"]["downloads"], 0)
        self.assertEqual(report["statistics"]["reused"], 3)
        self.assertEqual(self.download.call_count, 1)
        self.assertEqual(self.info.call_count, 2)
        self.assertEqual(len(self.store.load_documents(42, readable_only=True)), 3)

    def test_changed_timestamp_and_body_refresh_only_changed_documents(self):
        self.run_ingestion()
        self.metadata["updated_at"] = "2026-10-02T00:00:00Z"
        self.content = make_pdf(["Updated lecture"])
        self.snapshot["assignments"][0]["metadata"]["description"] = "<p>Changed instructions</p>"
        report = self.run_ingestion()
        self.assertEqual(report["statistics"]["downloads"], 1)
        self.assertEqual(report["statistics"]["reused"], 1)
        self.assertIn("Updated lecture", self.store.load_document(42, "file:9")["text"])
        self.assertEqual(self.store.load_document(42, "assignment:7")["text"], "Changed instructions")

    def test_permission_failure_keeps_other_documents_and_marks_old_pdf_stale(self):
        self.run_ingestion()
        response = Mock(status_code=403)
        self.info.side_effect = requests.HTTPError(response=response)
        report = self.run_ingestion()
        self.assertEqual(report["statistics"]["unavailable"], 1)
        self.assertEqual(report["statistics"]["readable"], 2)
        self.assertEqual(report["snapshot"]["coverage"], "partial")
        self.assertTrue(self.store.load_document(42, "file:9")["stale"])
        self.assertEqual(len(self.store.load_documents(42, readable_only=True)), 2)

    def test_listing_metadata_is_reused_without_a_second_metadata_request(self):
        self.snapshot["file_inventory"][0]["metadata"] = self.metadata
        self.run_ingestion()
        self.info.assert_not_called()
        self.download.assert_called_once()

    def test_missing_timestamps_require_download_but_skip_duplicate_parsing(self):
        self.metadata.pop("updated_at")
        self.run_ingestion()
        with patch("backend.app.document_service.extract_file") as extract:
            report = self.run_ingestion()
        extract.assert_not_called()
        self.assertEqual(report["statistics"]["downloads"], 1)
        self.assertEqual(report["statistics"]["reused"], 3)

    def test_unsupported_file_is_recorded_without_downloading(self):
        self.metadata["filename"] = "sheet.xlsx"
        self.metadata["display_name"] = "sheet.xlsx"
        report = self.run_ingestion()
        self.download.assert_not_called()
        self.assertEqual(report["statistics"]["unsupported"], 1)
        self.assertEqual(report["snapshot"]["coverage"], "partial")

    def test_newly_supported_format_retries_an_old_unsupported_result(self):
        self.metadata["filename"] = "assignment.ipynb"
        self.metadata["display_name"] = "assignment.ipynb"
        with patch("backend.app.document_service.file_format", return_value=None):
            self.run_ingestion()
        self.content = b'{"nbformat":4,"cells":[{"cell_type":"markdown","source":"New reader"}]}'
        report = self.run_ingestion()
        self.assertEqual(report["statistics"]["downloads"], 1)
        self.assertEqual(self.store.load_document(42, "file:9")["status"], "ok")

    def test_image_only_pdf_is_not_claimed_as_readable_and_not_reprocessed(self):
        self.content = make_pdf([None])
        report = self.run_ingestion()
        self.assertEqual(report["statistics"]["no_text"], 1)
        self.assertEqual(report["statistics"]["readable"], 2)
        report = self.run_ingestion()
        self.assertEqual(report["statistics"]["no_text"], 1)
        self.assertEqual(report["statistics"]["downloads"], 0)

    def test_force_ignores_cache_and_new_parser_version_invalidates_cache(self):
        self.run_ingestion()
        self.assertEqual(self.run_ingestion(force=True)["statistics"]["downloads"], 1)
        with patch("backend.app.document_service.PIPELINE_VERSION", "new-parser"):
            self.assertEqual(self.run_ingestion()["statistics"]["downloads"], 1)

    def test_file_limit_is_explicit_and_does_not_prevent_html_ingestion(self):
        report = self.run_ingestion(max_files=0)
        self.assertEqual(report["statistics"]["deferred"], 1)
        self.assertEqual(report["statistics"]["readable"], 2)
        self.info.assert_not_called()
        self.download.assert_not_called()

    def test_invalid_input_does_not_contact_canvas(self):
        for invalid in (0, -1, "42", True):
            with self.assertRaises(ValueError):
                ingestion_service.ingest_course(invalid, store=self.store)
        with self.assertRaises(ValueError):
            self.run_ingestion(max_files=-1)
        self.discover.assert_not_called()

    def test_locked_file_does_not_reuse_previously_readable_text(self):
        self.run_ingestion()
        self.metadata["locked_for_user"] = True
        report = self.run_ingestion()
        self.assertEqual(report["statistics"]["unavailable"], 1)
        self.assertTrue(self.store.load_document(42, "file:9")["stale"])

    def test_summary_is_concise_and_never_prints_download_urls_or_raw_text(self):
        report = self.run_ingestion()
        summary = ingestion_service.format_ingestion_summary(report)
        self.assertIn("Readable documents: 3", summary)
        self.assertNotIn(self.metadata["url"], summary)
        self.assertNotIn("Linear algebra lecture", summary)

    def test_office_files_use_existing_ingestion_storage_and_incremental_cache(self):
        import docx
        from backend.tests.test_office_documents import make_pptx, save_bytes

        word = docx.Document()
        word.add_paragraph("The course syllabus")
        for filename, content, units in (("lecture.pptx", make_pptx(), "slides"), ("syllabus.docx", save_bytes(word), "blocks")):
            with self.subTest(filename=filename):
                self.metadata.update(filename=filename, display_name=filename, size=len(content))
                self.content = content
                first = self.run_ingestion()
                self.assertEqual(first["statistics"]["downloads"], 1)
                saved = self.store.load_document(42, "file:9")
                self.assertEqual(saved["status"], "ok")
                self.assertTrue(saved[units])
                with patch("backend.app.document_service.extract_file") as extract:
                    second = self.run_ingestion()
                extract.assert_not_called()
                self.assertEqual(second["statistics"]["downloads"], 0)
                self.assertEqual(second["statistics"]["reused"], 3)

    def test_office_file_denial_preserves_other_course_content(self):
        self.metadata.update(filename="lecture.pptx", display_name="lecture.pptx")
        self.download.side_effect = requests.HTTPError(response=Mock(status_code=403))
        report = self.run_ingestion()
        self.assertEqual(report["statistics"]["unavailable"], 1)
        self.assertEqual(report["statistics"]["readable"], 2)
        self.assertEqual(report["snapshot"]["coverage"], "partial")

    def test_legacy_format_reprocesses_old_unsupported_cache_then_reuses_saved_text(self):
        import docx
        from backend.app import legacy_office
        from backend.tests.test_office_documents import save_bytes

        word = docx.Document()
        word.add_paragraph("Legacy assignment")
        self.metadata.update(filename="assignment.doc", display_name="assignment.doc")
        self.content = legacy_office.OLE_SIGNATURE + b"\x00" * 504
        with patch("backend.app.document_service.file_format", return_value=None):
            self.run_ingestion()
        with patch.object(legacy_office, "converter_fingerprint", return_value="test-converter"), patch.object(legacy_office, "convert_legacy_office", return_value=save_bytes(word)) as convert:
            first = self.run_ingestion()
            second = self.run_ingestion()
        self.assertEqual(first["statistics"]["downloads"], 1)
        self.assertEqual(second["statistics"]["downloads"], 0)
        convert.assert_called_once()
        saved = self.store.load_document(42, "file:9")
        self.assertEqual(saved["format"], "doc")
        self.assertIn("Legacy assignment", saved["text"])

    def test_missing_legacy_converter_does_not_block_other_course_documents(self):
        from backend.app import legacy_office

        self.metadata.update(filename="lecture.ppt", display_name="lecture.ppt")
        self.content = legacy_office.OLE_SIGNATURE + b"\x00" * 504
        with patch.object(legacy_office, "libreoffice_binary", return_value=None):
            report = self.run_ingestion()
        self.assertEqual(report["statistics"]["unavailable"], 1)
        self.assertEqual(report["statistics"]["readable"], 2)
        self.assertEqual(report["snapshot"]["coverage"], "partial")
