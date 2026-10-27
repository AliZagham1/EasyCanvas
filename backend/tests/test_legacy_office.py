import hashlib
import io
import os
import signal
import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

import docx
import pptx

from backend.app import document_service, legacy_office
from backend.tests.test_office_documents import make_pptx, save_bytes

LEGACY_BYTES = legacy_office.OLE_SIGNATURE + b"\x00" * 504


class LegacyOfficeTests(unittest.TestCase):
    def setUp(self):
        self.binary = patch.object(legacy_office, "libreoffice_binary", return_value="/fake/soffice").start()
        patch.object(legacy_office, "word_bridge_binary", return_value=None).start()
        patch.object(legacy_office, "converter_fingerprint", return_value="LibreOffice-test").start()
        self.addCleanup(patch.stopall)
        word = docx.Document()
        word.add_paragraph("Legacy assignment instructions")
        self.word_bytes = save_bytes(word)
        self.presentation_bytes = make_pptx()
        self.temporary_roots = []

    def converter(self, command, directory):
        root = Path(directory)
        self.temporary_roots.append(root)
        source = Path(command[-1])
        target = "docx" if source.suffix == ".doc" else "pptx"
        self.assertEqual(source.read_bytes(), LEGACY_BYTES)
        self.assertIn("--headless", command)
        profile = root / "profile/user/registrymodifications.xcu"
        self.assertIn('DisableMacrosExecution', profile.read_text())
        output = root / "output" / f"input.{target}"
        output.write_bytes(self.word_bytes if target == "docx" else self.presentation_bytes)
        return 0

    def test_detects_old_office_extensions_and_mime_types(self):
        for metadata, expected in (
            ({"filename": "assignment.DOC"}, "doc"),
            ({"display_name": "lecture.PPT"}, "ppt"),
            ({"content-type": "application/msword"}, "doc"),
            ({"content_type": "application/vnd.ms-powerpoint; charset=binary"}, "ppt"),
        ):
            self.assertEqual(document_service.file_format(metadata), expected)

    def test_doc_conversion_reuses_word_reader_and_hashes_original_bytes(self):
        with patch.object(legacy_office, "_run_conversion", side_effect=self.converter):
            result = document_service.extract_file({"filename": "assignment.doc"}, LEGACY_BYTES)
        self.assertEqual(result["status"], "partial")
        self.assertIn("legacy_doc_structure_unverified", result["warnings"])
        self.assertEqual(result["format"], "doc")
        self.assertEqual(result["converted_format"], "docx")
        self.assertIn("Legacy assignment instructions", result["text"])
        self.assertTrue(result["blocks"])
        self.assertEqual(result["content_hash"], hashlib.sha256(LEGACY_BYTES).hexdigest())
        self.assertFalse(any(path.exists() for path in self.temporary_roots))

    def test_ppt_conversion_preserves_slide_numbers_tables_and_notes(self):
        with patch.object(legacy_office, "_run_conversion", side_effect=self.converter):
            result = document_service.extract_file({"filename": "lecture.ppt"}, LEGACY_BYTES)
        self.assertEqual(result["status"], "ok")
        self.assertEqual(result["format"], "ppt")
        self.assertEqual(result["slide_count"], 1)
        self.assertEqual(result["slides"][0]["slide_number"], 1)
        self.assertIn("Discuss time slices.", result["text"])
        self.assertIn("Algorithm\tQueue", result["text"])
        self.assertFalse(any(path.exists() for path in self.temporary_roots))

    def test_missing_converter_is_explicit_and_does_not_crash(self):
        self.binary.return_value = None
        result = document_service.extract_file({"filename": "assignment.doc"}, LEGACY_BYTES)
        self.assertEqual(result["status"], "unavailable")
        self.assertEqual(result["error_type"], "LibreOfficeNotInstalled")

    def test_corrupt_input_never_starts_a_converter(self):
        with patch.object(legacy_office, "_run_conversion") as run:
            result = document_service.extract_file({"filename": "assignment.doc"}, b"not an Office document")
        self.assertEqual(result["status"], "error")
        self.assertEqual(result["error_type"], "InvalidLegacyOffice")
        run.assert_not_called()

    def test_converter_failure_and_missing_output_are_not_success(self):
        for code in (0, 1):
            roots = []
            def fail(command, directory):
                roots.append(Path(directory))
                return code
            with patch.object(legacy_office, "_run_conversion", side_effect=fail):
                result = document_service.extract_file({"filename": "assignment.doc"}, LEGACY_BYTES)
            self.assertEqual(result["error_type"], "LegacyConversionFailed")
            self.assertFalse(any(root.exists() for root in roots))

    def test_timeout_and_output_size_limit_clean_up_temporary_files(self):
        def timeout(command, directory):
            self.temporary_roots.append(Path(directory))
            raise legacy_office.LegacyConversionError("LegacyConversionTimeout")
        with patch.object(legacy_office, "_run_conversion", side_effect=timeout):
            result = document_service.extract_file({"filename": "lecture.ppt"}, LEGACY_BYTES)
        self.assertEqual(result["error_type"], "LegacyConversionTimeout")
        with patch.object(legacy_office, "MAX_LEGACY_BYTES", 1024), patch.object(legacy_office, "_run_conversion", side_effect=self.converter):
            result = document_service.extract_file({"filename": "assignment.doc"}, LEGACY_BYTES)
        self.assertEqual(result["status"], "too_large")
        self.assertFalse(any(path.exists() for path in self.temporary_roots))

    def test_conversion_process_does_not_receive_canvas_or_ai_secrets(self):
        with patch.dict(os.environ, {"CANVAS_TOKEN": "test-only-canvas", "OPENAI_API_KEY": "test-only-api"}):
            environment = legacy_office._converter_environment()
        self.assertNotIn("CANVAS_TOKEN", environment)
        self.assertNotIn("OPENAI_API_KEY", environment)

    @unittest.skipUnless(os.name == "posix", "POSIX process group behavior")
    def test_timeout_kills_only_its_own_process_group(self):
        process = Mock(pid=12345)
        process.wait.side_effect = [subprocess.TimeoutExpired("soffice", 60), 0]
        with patch.object(legacy_office.subprocess, "Popen", return_value=process), patch.object(legacy_office.os, "killpg") as kill:
            with self.assertRaises(legacy_office.LegacyConversionError):
                legacy_office._run_conversion(["soffice"], ".")
        kill.assert_called_once_with(12345, signal.SIGKILL)

    def test_converter_version_changes_legacy_cache_without_changing_modern_readers(self):
        before = document_service.file_pipeline_version({"filename": "lecture.ppt"})
        modern = document_service.file_pipeline_version({"filename": "lecture.pptx"})
        legacy_office.converter_fingerprint.return_value = "LibreOffice-new"
        self.assertNotEqual(before, document_service.file_pipeline_version({"filename": "lecture.ppt"}))
        self.assertEqual(modern, document_service.file_pipeline_version({"filename": "lecture.pptx"}))

    def test_native_word_table_bridge_is_temporary_and_uses_the_existing_reader(self):
        commands = []
        def run(command, directory):
            commands.append(command)
            root = Path(directory)
            self.temporary_roots.append(root)
            if command[0] == "/fake/textutil":
                (root / "input.html").write_text("<table><tr><td>Labs</td><td>30%</td></tr></table>")
            else:
                self.assertEqual(Path(command[-1]).suffix, ".html")
                (root / "output/input.docx").write_bytes(self.word_bytes)
            return 0
        with patch.object(legacy_office, "word_bridge_binary", return_value="/fake/textutil"), patch.object(legacy_office, "_run_conversion", side_effect=run):
            result = document_service.extract_file({"filename": "assignment.doc"}, LEGACY_BYTES)
        self.assertEqual(len(commands), 2)
        self.assertIn("-noload", commands[0])
        self.assertIn("Legacy assignment instructions", result["text"])
        self.assertFalse(any(root.exists() for root in self.temporary_roots))


@unittest.skipUnless(os.getenv("EASYCANVAS_TEST_LEGACY_CONVERSION") == "1", "Set EASYCANVAS_TEST_LEGACY_CONVERSION=1 for real LibreOffice conversions")
class LegacyOfficeIntegrationTests(unittest.TestCase):
    def make_legacy(self, content, source_format, target_format, filter_name, root):
        binary = legacy_office.libreoffice_binary()
        self.assertIsNotNone(binary, "LibreOffice must be installed for integration tests")
        source = root / f"fixture.{source_format}"
        source.write_bytes(content)
        profile = root / "seed-profile"
        legacy_office._prepare_profile(profile)
        code = legacy_office._run_conversion([
            binary, f"-env:UserInstallation={profile.as_uri()}",
            "--headless", "--nologo", "--nodefault", "--norestore",
            "--convert-to", f"{target_format}:{filter_name}", "--outdir", str(root), str(source),
        ], root)
        self.assertEqual(code, 0)
        legacy = (root / f"fixture.{target_format}").read_bytes()
        self.assertTrue(legacy.startswith(legacy_office.OLE_SIGNATURE), "Fixture must be a real binary Office file")
        return legacy

    def test_real_binary_doc_conversion_preserves_paragraphs_and_table_text(self):
        word = docx.Document()
        word.add_paragraph("Legacy syllabus policy")
        table = word.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Labs"
        table.cell(0, 1).text = "30%"
        word.add_paragraph("After the table")
        with tempfile.TemporaryDirectory(prefix="easycanvas-legacy-test-") as directory:
            legacy = self.make_legacy(save_bytes(word), "docx", "doc", "MS Word 97", Path(directory).resolve())
            result = document_service.extract_file({"filename": "syllabus.doc"}, legacy)
        self.assertIn(result["status"], ("ok", "partial"))
        self.assertIn("Legacy syllabus policy", result["text"])
        self.assertIn("Labs\t30%", result["text"])
        self.assertEqual(result["format"], "doc")
        self.assertEqual(result["content_hash"], hashlib.sha256(legacy).hexdigest())

    def test_binary_doc_table_at_end_is_not_claimed_as_verified_structure(self):
        word = docx.Document()
        word.add_paragraph("Before the table")
        table = word.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Labs"
        table.cell(0, 1).text = "30%"
        with tempfile.TemporaryDirectory(prefix="easycanvas-legacy-test-") as directory:
            legacy = self.make_legacy(save_bytes(word), "docx", "doc", "MS Word 97", Path(directory).resolve())
            result = document_service.extract_file({"filename": "syllabus.doc"}, legacy)
        self.assertEqual(result["status"], "partial")
        self.assertIn("legacy_doc_structure_unverified", result["warnings"])
        self.assertIn("Labs", result["text"])
        self.assertIn("30%", result["text"])
        if legacy_office.word_bridge_binary():
            self.assertIn("Labs\t30%", result["text"])

    def test_real_binary_ppt_conversion_preserves_slides_tables_and_speaker_notes(self):
        presentation = pptx.Presentation(io.BytesIO(make_pptx()))
        presentation.slides.add_slide(presentation.slide_layouts[6])
        with tempfile.TemporaryDirectory(prefix="easycanvas-legacy-test-") as directory:
            legacy = self.make_legacy(save_bytes(presentation), "pptx", "ppt", "MS PowerPoint 97", Path(directory).resolve())
            result = document_service.extract_file({"filename": "lecture.ppt"}, legacy)
        self.assertIn(result["status"], ("ok", "partial"))
        self.assertEqual(result["format"], "ppt")
        self.assertEqual(result["slide_count"], 2)
        self.assertEqual([slide["slide_number"] for slide in result["slides"]], [1, 2])
        self.assertIn("Process scheduling", result["text"])
        self.assertIn("Discuss time slices.", result["text"])
        self.assertIn("Algorithm\tQueue", result["text"])


if __name__ == "__main__":
    unittest.main()
