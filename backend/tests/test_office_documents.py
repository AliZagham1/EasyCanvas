import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import docx
import pptx
from PIL import Image
from pptx.util import Inches

from backend.app import document_service
from backend.app.ingestion_store import IngestionStore


def save_bytes(document):
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def image_bytes():
    output = io.BytesIO()
    Image.new("RGB", (2, 2), color="white").save(output, format="PNG")
    output.seek(0)
    return output


def make_pptx():
    """Real Office fixture containing visible text, grouped text, a table, and notes."""
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Operating systems"
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text = "Process scheduling"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
    for row, values in enumerate((("Algorithm", "Queue"), ("Round robin", "Ready"))):
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    slide.notes_slide.notes_text_frame.text = "Discuss time slices."
    return save_bytes(presentation)


class OfficeDocumentTests(unittest.TestCase):
    def test_detects_extensions_and_mime_types_including_generic_canvas_file_names(self):
        cases = (
            ({"filename": "SYLLABUS.DOCX", "content-type": "application/octet-stream"}, "docx"),
            ({"display_name": "lecture.PPTX"}, "pptx"),
            ({"filename": "download", "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document; charset=binary"}, "docx"),
            ({"content-type": "application/vnd.openxmlformats-officedocument.presentationml.presentation"}, "pptx"),
        )
        for metadata, expected in cases:
            with self.subTest(expected=expected):
                self.assertEqual(document_service.file_format(metadata), expected)

    def test_docx_preserves_paragraph_table_order_and_does_not_invent_page_numbers(self):
        document = docx.Document()
        document.add_heading("Syllabus", level=1)
        document.add_paragraph("Before the table")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Category"
        table.cell(0, 1).text = "Weight"
        table.cell(1, 0).text = "Labs"
        table.cell(1, 1).text = "30%"
        document.add_paragraph("After the table")
        result = document_service.extract_file({"filename": "syllabus.docx"}, save_bytes(document))
        self.assertEqual(result["status"], "ok")
        self.assertLess(result["text"].index("Before the table"), result["text"].index("Category\tWeight"))
        self.assertLess(result["text"].index("Labs\t30%"), result["text"].index("After the table"))
        self.assertEqual([block["kind"] for block in result["blocks"]], ["paragraph", "paragraph", "table", "paragraph"])
        self.assertEqual(result["blocks"][2]["block_number"], 3)
        self.assertIsNone(result["page_count"])

    def test_docx_reads_nested_tables_merged_cells_headers_and_footers(self):
        document = docx.Document()
        document.sections[0].header.paragraphs[0].text = "Course handbook"
        document.sections[0].footer.paragraphs[0].text = "Instructor contact"
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged heading"
        nested = table.cell(0, 0).add_table(rows=1, cols=1)
        nested.cell(0, 0).text = "Nested detail"
        result = document_service.extract_docx(save_bytes(document))
        self.assertEqual(result["status"], "ok")
        self.assertEqual(result["text"].count("Merged heading"), 1)
        self.assertIn("Nested detail", result["text"])
        self.assertIn("Course handbook", result["text"])
        self.assertIn("Instructor contact", result["text"])
        self.assertTrue(any(block["location"] == "section 1 header" for block in result["blocks"]))

    def test_docx_images_are_flagged_instead_of_claimed_as_read(self):
        document = docx.Document()
        document.add_paragraph("Read the figure")
        document.add_picture(image_bytes())
        result = document_service.extract_docx(save_bytes(document))
        self.assertEqual(result["status"], "partial")
        self.assertIn("images_not_extracted", result["warnings"])
        self.assertIn("Read the figure", result["text"])

    def test_pptx_reads_grouped_text_tables_and_speaker_notes(self):
        result = document_service.extract_file({"filename": "lecture.pptx"}, make_pptx())
        self.assertEqual(result["status"], "ok")
        self.assertEqual(result["slide_count"], 1)
        slide = result["slides"][0]
        self.assertEqual(slide["slide_number"], 1)
        self.assertEqual(slide["title"], "Operating systems")
        self.assertIn("Process scheduling", slide["text"])
        self.assertIn("Algorithm\tQueue", slide["text"])
        self.assertIn("Round robin\tReady", slide["text"])
        self.assertEqual(slide["text"][slide["notes_start_character"]:], "Discuss time slices.")
        self.assertIsNone(result["page_count"])

    def test_pptx_keeps_blank_and_image_only_slide_numbers(self):
        presentation = pptx.Presentation()
        first = presentation.slides.add_slide(presentation.slide_layouts[5])
        first.shapes.title.text = "Readable slide"
        presentation.slides.add_slide(presentation.slide_layouts[6])
        third = presentation.slides.add_slide(presentation.slide_layouts[6])
        third.shapes.add_picture(image_bytes(), Inches(1), Inches(1))
        result = document_service.extract_pptx(save_bytes(presentation))
        self.assertEqual(result["status"], "partial")
        self.assertEqual(result["slide_count"], 3)
        self.assertEqual([slide["slide_number"] for slide in result["slides"]], [1, 2, 3])
        self.assertEqual(result["slides"][1]["status"], "no_text")
        self.assertEqual(result["slides"][2]["status"], "no_text")
        self.assertIn("images_not_extracted", result["slides"][2]["warnings"])

    def test_empty_documents_and_corrupt_or_wrong_office_packages_are_explicit(self):
        self.assertEqual(document_service.extract_docx(save_bytes(docx.Document()))["status"], "no_text")
        self.assertEqual(document_service.extract_pptx(save_bytes(pptx.Presentation()))["status"], "no_text")
        for reader in (document_service.extract_docx, document_service.extract_pptx):
            self.assertEqual(reader(b"not a zip file")["status"], "error")
        result = document_service.extract_docx(make_pptx())
        self.assertEqual(result["error_type"], "WrongOfficeFormat")

    def test_archive_and_slide_limits_stop_work_before_expensive_parsing(self):
        content = make_pptx()
        with patch.object(document_service, "MAX_OFFICE_UNCOMPRESSED_BYTES", 1), patch("pptx.Presentation") as parser:
            self.assertEqual(document_service.extract_pptx(content)["status"], "too_large")
            parser.assert_not_called()
        with patch.object(document_service, "MAX_PPTX_SLIDES", 0):
            self.assertEqual(document_service.extract_pptx(content)["status"], "too_large")

    def test_office_parser_versions_do_not_invalidate_unrelated_pdf_caches(self):
        before = document_service.file_pipeline_version({"filename": "lecture.pptx"})
        pdf_version = document_service.file_pipeline_version({"filename": "lecture.pdf"})
        with patch.object(pptx, "__version__", "next-version"):
            self.assertNotEqual(before, document_service.file_pipeline_version({"filename": "lecture.pptx"}))
            self.assertEqual(pdf_version, document_service.file_pipeline_version({"filename": "lecture.pdf"}))

    def test_storage_keeps_slide_and_word_block_metadata_without_duplicating_text(self):
        word = docx.Document()
        word.add_paragraph("Syllabus policy")
        for format_name, content in (("pptx", make_pptx()), ("docx", save_bytes(word))):
            with self.subTest(format=format_name), tempfile.TemporaryDirectory() as directory:
                store = IngestionStore("test-student", Path(directory) / "ingestion.sqlite3")
                result = document_service.extract_file({"filename": "file." + format_name}, content)
                document = {**result, "document_key": "file:9", "version": "v1"}
                store.save_document(42, document, "first")
                loaded = IngestionStore("test-student", store.path).load_document(42, "file:9")
                self.assertEqual(loaded["text"], result["text"])
                unit = "slides" if format_name == "pptx" else "blocks"
                self.assertEqual(loaded[unit], result[unit])
                self.assertNotIn(unit, loaded["latest_attempt"])


if __name__ == "__main__":
    unittest.main()
